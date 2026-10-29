"""Build the adviser-facing research proposal from reviewed Markdown."""
from pathlib import Path
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "AI招聘低分成因诊断_研究计划_20261027.docx"


def font(style, size, chinese="宋体", latin="Calibri"):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    properties = style.element.get_or_add_rPr()
    names = properties.find(qn("w:rFonts"))
    if names is None:
        names = OxmlElement("w:rFonts")
        properties.insert(0, names)
    for key in ("ascii", "hAnsi", "cs"):
        names.set(qn("w:" + key), latin)
    names.set(qn("w:eastAsia"), chinese)
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        names.attrib.pop(qn("w:" + key), None)
    language = properties.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        properties.append(language)
    language.set(qn("w:eastAsia"), "zh-CN")


def main():
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    font(normal, 11)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True
    font(document.styles["Title"], 17, "黑体")
    document.styles["Title"].paragraph_format.space_after = Pt(5)
    document.styles["Title"].paragraph_format.keep_with_next = True
    font(document.styles["Subtitle"], 10.5)
    document.styles["Subtitle"].paragraph_format.space_after = Pt(8)
    document.styles["Subtitle"].font.italic = False
    font(document.styles["Heading 1"], 12, "黑体")
    document.styles["Heading 1"].paragraph_format.space_before = Pt(6)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(3)
    document.styles["Heading 1"].paragraph_format.keep_with_next = True
    reference_style = document.styles.add_style("Proposal Reference", 1)
    reference_style.base_style = normal
    font(reference_style, 10.5)
    reference_style.paragraph_format.line_spacing = 1.0
    reference_style.paragraph_format.space_after = Pt(2)
    for style in document.styles:
        properties = style.element.find(qn("w:pPr"))
        if properties is not None:
            for element in list(properties):
                if element.tag in (qn("w:pBdr"), qn("w:shd")):
                    properties.remove(element)
    source = (ROOT / "01_研究计划.md").read_text(encoding="utf-8")
    for block in re.split(r"\n\s*\n", source.strip()):
        text = block.replace("\n", " ").strip()
        if text == "<!-- PAGE_BREAK -->":
            document.add_page_break()
        elif text.startswith("# "):
            document.add_paragraph(text[2:], "Title")
        elif text.startswith("## "):
            document.add_paragraph(text[3:], "Heading 1")
        elif text.startswith("研究计划 "):
            document.add_paragraph(text, "Subtitle")
        elif re.match(r"^\[\d+\]", text):
            document.add_paragraph(text, "Proposal Reference")
        else:
            paragraph = document.add_paragraph(text)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    # Clear theme-linked borders/colours on title/heading paragraphs.
    for paragraph in document.paragraphs:
        if paragraph.style.name in ("Title", "Subtitle", "Heading 1"):
            properties = paragraph._p.get_or_add_pPr()
            for element in list(properties):
                if element.tag in (qn("w:pBdr"), qn("w:shd")):
                    properties.remove(element)
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.font.underline = False
    document.core_properties.title = "AI招聘中的低分成因诊断与补救机会配置"
    document.core_properties.subject = "Research proposal"
    document.core_properties.author = ""
    document.core_properties.comments = ""
    document.save(OUTPUT)
    print(str(OUTPUT))


if __name__ == "__main__":
    main()
